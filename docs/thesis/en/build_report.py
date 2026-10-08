"""Build the graduation report from the faculty template (template/report-template.docx).

  python build_report.py              -> docs/report.docx    from docs/thesis/en/*.md (English)
  python build_report.py --lang vi    -> docs/report-vi.docx from docs/thesis/vi/*.md (Vietnamese,
                                         figures from build/png-vi, numbers as 1.234,5)
Then run finalize.ps1 -Docx <file> to fill the table of contents and export a PDF with Word.

Markup understood in the chapter files:
  # Title / ## Title / ### Title     -> Heading 1/2/3 (numbered by the template: CHAPTER n., n.m, n.m.k)
  plain paragraph                     -> body text (style "Nội dung văn bản")
  - item / 1. item                    -> bullet / numbered list
  TABLE: caption  + markdown table    -> caption above, table
  FIG: file.png | caption | widthCm   -> picture, caption below
  CODE: caption + ```fenced block```  -> listing caption above, monospace block
  **bold**, *italic*, `code`, [@key] -> inline formatting, IEEE citation numbers
  @ABSTRACT / @ABBREVIATIONS (front file only)
"""
import copy
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
DOCS = HERE.parent.parent
# Pristine copy of the faculty template; the filled report is written to docs/report.docx
# (English) or docs/report-vi.docx (python build_report.py --lang vi).
TEMPLATE = HERE / "template" / "report-template.docx"
_argv = sys.argv[1:]
LANG = "en"
if "--lang" in _argv:
    _i = _argv.index("--lang")
    LANG = _argv[_i + 1]
    del _argv[_i:_i + 2]
assert LANG in ("en", "vi"), "--lang must be en or vi"
_args = [a for a in _argv if not a.startswith("--")]
CONTENT = HERE if LANG == "en" else HERE.parent / "vi"
OUT = Path(_args[0]) if _args else DOCS / ("report.docx" if LANG == "en" else "report-vi.docx")
PNG = HERE / "build" / ("png" if LANG == "en" else "png-vi")

AUTHOR = "VO ANH KIET" if LANG == "en" else "VÕ ANH KIỆT"
STUDENT_ID = "523H0048"
TITLE = {
    "en": "DEVELOPMENT OF A REAL-TIME COLLABORATIVE WORKSPACE PLATFORM WITH INTEGRATED MULTIPOINT VIDEO CONFERENCING",
    "vi": "PHÁT TRIỂN NỀN TẢNG KHÔNG GIAN LÀM VIỆC CỘNG TÁC THỜI GIAN THỰC TÍCH HỢP HỘI NGHỊ TRUYỀN HÌNH ĐA ĐIỂM",
}[LANG]
DATE_LINE = {"en": "Ho Chi Minh City, October 2026", "vi": "TP. Hồ Chí Minh, tháng 10 năm 2026"}[LANG]
T = {
    "en": dict(thesis="GRADUATION THESIS", signature="(Signature and full name)", signer="Vo Anh Kiet",
               abstract="ABSTRACT", abbr="LIST OF ABBREVIATIONS", figures="LIST OF FIGURES", tables="LIST OF TABLES",
               refs="REFERENCES", chapter="CHAPTER %1. ", caption_words="Figure|Table|Listing", missing="missing image"),
    "vi": dict(thesis="ĐỒ ÁN TỐT NGHIỆP", signature="(Ký và ghi rõ họ tên)", signer="Võ Anh Kiệt",
               abstract="TÓM TẮT", abbr="DANH MỤC CÁC CHỮ VIẾT TẮT", figures="DANH MỤC HÌNH VẼ", tables="DANH MỤC BẢNG BIỂU",
               refs="TÀI LIỆU THAM KHẢO", chapter="CHƯƠNG %1. ", caption_words="Hình|Bảng|Đoạn mã", missing="thiếu hình"),
}[LANG]
ACK_VI = (
    "Em xin gửi lời cảm ơn chân thành đến thầy Vũ Đình Hồng, người đã trực tiếp hướng dẫn em trong suốt quá "
    "trình thực hiện đồ án tốt nghiệp. Những góp ý quý báu, sự kiên nhẫn và các nhận xét cụ thể của thầy về thiết "
    "kế hệ thống, phương pháp thực nghiệm và cách trình bày báo cáo đã giúp em hoàn thành đồ án này.",
    "Em cũng xin cảm ơn quý thầy cô Khoa Công nghệ Thông tin, Trường Đại học Tôn Đức Thắng đã truyền đạt kiến thức "
    "trong suốt những năm học vừa qua, cùng gia đình và bạn bè đã động viên và giúp em thử nghiệm ứng dụng. Dù đã "
    "cố gắng, đồ án chắc chắn vẫn còn những thiếu sót; em rất mong nhận được ý kiến đóng góp của quý thầy cô.",
)
DECL_VI = [
    "Em xin cam đoan đây là đồ án tốt nghiệp của riêng em, được thực hiện dưới sự hướng dẫn của thầy Vũ Đình Hồng. "
    "Các nội dung nghiên cứu và kết quả trong đồ án là trung thực và chưa được công bố dưới bất kỳ hình thức nào "
    "trước đây. Số liệu trong các bảng biểu và hình vẽ phục vụ phân tích, đánh giá do chính em thu thập từ các thực "
    "nghiệm trình bày ở Chương 5, có ghi rõ nguồn gốc trong báo cáo.",
    "Ngoài ra, đồ án có sử dụng một số ý tưởng, đặc tả và số liệu của các tác giả, tổ chức khác; tất cả đều được "
    "trích dẫn và liệt kê trong phần tài liệu tham khảo.",
    "Nếu phát hiện có bất kỳ gian lận nào, em xin hoàn toàn chịu trách nhiệm về nội dung đồ án của mình. Trường "
    "Đại học Tôn Đức Thắng không liên quan đến những vi phạm tác quyền, bản quyền do em gây ra trong quá trình "
    "thực hiện (nếu có).",
]
# Fixed texts of the template, replaced in the Vietnamese edition (index in the original template).
TEMPLATE_VI = {
    0: "TỔNG LIÊN ĐOÀN LAO ĐỘNG VIỆT NAM", 29: "TỔNG LIÊN ĐOÀN LAO ĐỘNG VIỆT NAM",
    1: "TRƯỜNG ĐẠI HỌC TÔN ĐỨC THẮNG", 30: "TRƯỜNG ĐẠI HỌC TÔN ĐỨC THẮNG",
    2: "KHOA CÔNG NGHỆ THÔNG TIN", 31: "KHOA CÔNG NGHỆ THÔNG TIN",
    28: "THÀNH PHỐ HỒ CHÍ MINH, NĂM 2026", 50: "THÀNH PHỐ HỒ CHÍ MINH, NĂM 2026",
    47: "Người hướng dẫn", 48: "Vũ Đình Hồng",
    51: "LỜI CẢM ƠN", 61: "LỜI CAM ĐOAN", 56: "Tác giả", 66: "Tác giả",
}
TEXT_WIDTH_DXA = 11906 - 1985 - 1134  # A4 minus template margins

BODY = "Nội dung văn bản"
warnings = []

# ---------------------------------------------------------------- references
REFS = {}
ref_order = []


def load_refs():
    for line in (HERE / "references.txt").read_text(encoding="utf-8").splitlines():
        if line.strip() and not line.startswith("#"):
            key, entry = line.split("|", 1)
            REFS[key.strip()] = entry.strip()


def cite(key):
    if key not in REFS:
        warnings.append(f"unknown reference {key}")
        return "[?]"
    if key not in ref_order:
        ref_order.append(key)
    return f"[{ref_order.index(key) + 1}]"


# ---------------------------------------------------------------- xml helpers
def set_text(p, value):
    """Replace a paragraph's text but keep the first run's formatting and non-run nodes."""
    runs = p.runs
    if not runs:
        p.add_run(value)
        return
    runs[0].text = value
    for r in runs[1:]:
        r._element.getparent().remove(r._element)


def add_run(p, text, bold=False, italic=False, mono=False, size=None):
    r = p.add_run(text)
    if bold:
        r.bold = True
    if italic:
        r.italic = True
    if mono:
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
        r.font.size = Pt(size or 11)
    elif size:
        r.font.size = Pt(size)
    return r


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[@[\w-]+\](?:,\s*\[@[\w-]+\])*)")


def add_inline(p, text, size=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            add_run(p, part[2:-2], bold=True, size=size)
        elif part.startswith("`"):
            add_run(p, part[1:-1], mono=True, size=(size - 1.5) if size else 11.5)
        elif part.startswith("[@"):
            keys = re.findall(r"\[@([\w-]+)\]", part)
            add_run(p, ", ".join(cite(k) for k in keys), size=size)
        elif part.startswith("*"):
            add_run(p, part[1:-1], italic=True, size=size)
        else:
            add_run(p, part, size=size)


def p_after(anchor, style=None):
    """Create a paragraph right after the given element and return it."""
    new = OxmlElement("w:p")
    anchor.addnext(new)
    para = docx.text.paragraph.Paragraph(new, doc._body)
    if style:
        para.style = doc.styles[style]
    return para


class Cursor:
    """Appends block elements after a moving anchor element."""

    def __init__(self, anchor):
        self.anchor = anchor

    def para(self, style=BODY):
        p = p_after(self.anchor, style)
        self.anchor = p._element
        return p

    def element(self, el):
        self.anchor.addnext(el)
        self.anchor = el


def set_ind(p, left=None, hanging=None, first=None):
    pPr = p._element.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    for k in ("w:left", "w:hanging", "w:firstLine"):
        if ind.get(qn(k)) is not None:
            del ind.attrib[qn(k)]
    if left is not None:
        ind.set(qn("w:left"), str(left))
    if hanging is not None:
        ind.set(qn("w:hanging"), str(hanging))
    if first is not None:
        ind.set(qn("w:firstLine"), str(first))


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def add_field(p, instr, placeholder="Right-click and choose Update Field."):
    def fld(kind):
        r = OxmlElement("w:r")
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), kind)
        r.append(f)
        return r

    p._element.append(fld("begin"))
    r = OxmlElement("w:r")
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = f" {instr} "
    r.append(t)
    p._element.append(r)
    p._element.append(fld("separate"))
    p.add_run(placeholder)
    p._element.append(fld("end"))


# ---------------------------------------------------------------- numbering
def setup_numbering():
    numbering = doc.part.numbering_part.element
    ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    nums = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    base_abs, base_num = max(ids) + 1, max(nums) + 1

    def abstract(aid, fmt, text_, font=None):
        a = OxmlElement("w:abstractNum")
        a.set(qn("w:abstractNumId"), str(aid))
        mlt = OxmlElement("w:multiLevelType")
        mlt.set(qn("w:val"), "hybridMultilevel")
        a.append(mlt)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        for tag, val in (("w:start", "1"), ("w:numFmt", fmt), ("w:lvlText", text_), ("w:lvlJc", "left")):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            lvl.append(e)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "1077")
        ind.set(qn("w:hanging"), "357")
        ppr.append(ind)
        lvl.append(ppr)
        if font:
            rpr = OxmlElement("w:rPr")
            rf = OxmlElement("w:rFonts")
            for k in ("w:ascii", "w:hAnsi", "w:hint"):
                rf.set(qn(k), font if k != "w:hint" else "default")
            rpr.append(rf)
            lvl.append(rpr)
        a.append(lvl)
        # abstractNum elements must precede num elements
        first_num = numbering.find(qn("w:num"))
        first_num.addprevious(a)

    abstract(base_abs, "bullet", "•", "Times New Roman")
    abstract(base_abs + 1, "decimal", "%1.")
    state = {"next": base_num}

    def new_num(kind):
        n = OxmlElement("w:num")
        nid = state["next"]
        state["next"] += 1
        n.set(qn("w:numId"), str(nid))
        an = OxmlElement("w:abstractNumId")
        an.set(qn("w:val"), str(base_abs if kind == "bullet" else base_abs + 1))
        n.append(an)
        if kind == "decimal":
            ov = OxmlElement("w:lvlOverride")
            ov.set(qn("w:ilvl"), "0")
            so = OxmlElement("w:startOverride")
            so.set(qn("w:val"), "1")
            ov.append(so)
            n.append(ov)
        numbering.append(n)
        return nid

    return new_num


def number_para(p, num_id):
    pPr = p._element.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl")
    il.set(qn("w:val"), "0")
    ni = OxmlElement("w:numId")
    ni.set(qn("w:val"), str(num_id))
    numPr.append(il)
    numPr.append(ni)
    style = pPr.find(qn("w:pStyle"))
    (style if style is not None else pPr[0]).addnext(numPr) if len(pPr) else pPr.append(numPr)
    set_ind(p, left=1077, hanging=357)


# ---------------------------------------------------------------- styles
def ensure_styles():
    st = doc.styles
    for name in ("Figure Caption", "Table Caption", "Listing Caption"):
        if name not in [s.name for s in st]:
            s = st.add_style(name, 1)
            s.base_style = st["Caption"]
            s.font.size = Pt(12)
            s.font.italic = name != "Table Caption"
            s.font.bold = name == "Table Caption"
            s.paragraph_format.space_before = Pt(4)
            s.paragraph_format.space_after = Pt(10)
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if name != "Figure Caption":
                s.paragraph_format.keep_with_next = True
    if "Code Block" not in [s.name for s in st]:
        s = st.add_style("Code Block", 1)
        s.base_style = st["Normal"]
        s.font.name = "Consolas"
        s.font.size = Pt(9)
        pf = s.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.3)
        pf.keep_with_next = True
        rpr = s.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        rf.set(qn("w:hAnsi"), "Consolas")
        rf.set(qn("w:cs"), "Consolas")
        shade(s.element.get_or_add_pPr(), "F3F5F7")
    # Entries of the lists of figures (toc 9) and tables (toc 8): plain, left-aligned, hanging indent
    for name in ("toc 8", "toc 9"):
        s = st[name]
        s.font.bold = False
        s.font.italic = False
        s.font.size = Pt(12)
        s.font.name = "Times New Roman"
        rf = s.element.get_or_add_rPr().find(qn("w:rFonts"))
        for k in ("w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(k), "Times New Roman")
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rf.get(qn(k)) is not None:
                del rf.attrib[qn(k)]
        pf = s.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent = Cm(1.6)
        pf.first_line_indent = Cm(-1.6)
        pf.space_before = Pt(0)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.15
    # Headings: start chapters on a new page
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(12)
    for h in ("Heading 2", "Heading 3"):
        doc.styles[h].paragraph_format.space_before = Pt(10)
        doc.styles[h].paragraph_format.space_after = Pt(4)
        doc.styles[h].paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- blocks
def table_block(cur, caption, rows, compact=False):
    cap = cur.para("Table Caption")
    add_inline(cap, caption)
    header, body = rows[0], rows[1:]
    ncol = len(header)
    # column widths proportional to content length, bounded
    # wide tables (measurement results) use a smaller font so "0.19 ± 0.10" fits on one line
    fsize = 10 if ncol >= 8 else 12
    body_lens = [max([max(len(x) for x in re.sub(r"[`*]", "", r[i]).split("<br>")) for r in rows[1:] if i < len(r)] or [0])
                 for i in range(ncol)]
    head_lens = [min(len(header[i]), 14) for i in range(ncol)]
    lens = [max(b_, h_) for b_, h_ in zip(body_lens, head_lens)]
    weights = [min(max(l, 2 if header[i].strip() == "N" else 6), 60) for i, l in enumerate(lens)]
    total = sum(weights)
    widths = [int(TEXT_WIDTH_DXA * w / total) for w in weights]
    widths[-1] += TEXT_WIDTH_DXA - sum(widths)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = doc.styles["Table Grid"]
    tbl = t._tbl
    tblPr = tbl.tblPr
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(TEXT_WIDTH_DXA))
    tw.set(qn("w:type"), "dxa")
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    tblPr.append(tw)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for ri, row in enumerate(rows):
        tr = t.rows[ri]
        if ri == 0:
            trPr = tr._tr.get_or_add_trPr()
            h = OxmlElement("w:tblHeader")
            trPr.append(h)
        for ci in range(ncol):
            cell = tr.cells[ci]
            cell.width = docx.shared.Twips(widths[ci])
            tcPr = cell._tc.get_or_add_tcPr()
            if ri == 0:
                shade(tcPr, "E7EBF0")
            p = cell.paragraphs[0]
            p.style = doc.styles["Normal"]
            pf = p.paragraph_format
            pf.space_before = Pt(0 if compact else 2)
            pf.space_after = Pt(0 if compact else 2)
            pf.line_spacing = 1.0 if compact else 1.1
            val = row[ci] if ci < len(row) else ""
            centered = val.strip() in ("✓", "–", "✓*") or re.fullmatch(r"[\d.,−+ /±]+", val.replace("<br>", " ").strip() or "x")
            if centered:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if ri == 0:
                add_run(p, re.sub(r"\*\*", "", val), bold=True, size=fsize)
            else:
                for k, part in enumerate(val.split("<br>")):
                    if k:
                        p.add_run().add_break()
                    add_inline(p, part, size=fsize)
    cur.element(tbl)
    spacer = cur.para("Normal")
    spacer.paragraph_format.space_after = Pt(2)
    return caption


def figure_block(cur, file, caption, width_cm):
    path = PNG / file
    p = cur.para("Hình")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    # single spacing: a 1.5 line multiple would also stretch the inline picture's line
    p.paragraph_format.line_spacing = 1.0
    if path.exists():
        p.add_run().add_picture(str(path), width=Cm(min(float(width_cm), 15.5)))
    else:
        warnings.append(f"missing image {file}")
        add_run(p, f"[{T['missing']} {file}]", italic=True)
    cap = cur.para("Figure Caption")
    add_inline(cap, caption)


def code_block(cur, caption, lines):
    cap = cur.para("Listing Caption")
    add_inline(cap, caption)
    for i, line in enumerate(lines):
        p = cur.para("Code Block")
        if i == len(lines) - 1:
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.space_after = Pt(10)
        if i == 0:
            p.paragraph_format.space_before = Pt(2)
        r = p.add_run(line if line else " ")
        r.font.name = "Consolas"


# ---------------------------------------------------------------- measured data
SUMMARY = None
DIGITS = {"upKbps": 0, "upP95": 0, "downKbps": 0, "cpu": 1, "encoders": 1, "rtt": 2, "audioJitter": 2,
          "videoJitter": 2, "loss": 2, "fps": 1, "jitterBuffer": 0, "nackPerMin": 1, "firstFrame": 0,
          "limited": 2, "workerCpu": 1, "nodeCpu": 1, "rss": 0, "freezePerClient": 1}


def row(label, n):
    for r in SUMMARY["labels"].get(label, {}).get("rows", []):
        if r["clients"] == n:
            return r
    warnings.append(f"no data for {label} N={n}")
    return None


def val(label, n, metric):
    r = row(label, n)
    return r[metric]["mean"] if r and r.get(metric) else float("nan")


def num(x, d=0):
    if x != x:  # NaN
        return "n/a"
    if abs(x) >= 100 and d > 0:
        d = 0 if abs(x) >= 1000 else min(d, 1)
    out = f"{x:,.{d}f}"
    return out if LANG == "en" else out.translate(str.maketrans({",": ".", ".": ","}))


def cell(label, n, metric, with_sd=True, br=False):
    r = row(label, n)
    if not r or not r.get(metric):
        return "–"
    s, d = r[metric], DIGITS.get(metric, 1)
    if with_sd and s.get("sd") is not None:
        return f"{num(s['mean'], d)}{'<br>' if br else ' '}± {num(s['sd'], d)}"
    return num(s["mean"], d)


LABELS = {"mesh": "mesh-norender", "sfu": "sfu-norender"}


def substitute(text):
    def token(m):
        only_mean, label, n, metric = m.group(1) == "=", m.group(2), int(m.group(3)), m.group(4)
        return cell(LABELS.get(label, label), n, metric, with_sd=not only_mean)

    text = re.sub(r"\{\{(=?)([\w-]+):(\d+):(\w+)\}\}", token, text)
    m, s = "mesh-norender", "sfu-norender"
    per_conn = val(m, 4, "upKbps") / 3
    computed = {
        "MESH_N3": num(2 * val(m, 2, "upKbps")),
        "UP_SAVING_4": num(100 * (1 - val(s, 4, "upKbps") / val(m, 4, "upKbps")), 0),
        "MESH_N6_THEORY": num(5 * per_conn, 0),
        "MESH_N8_THEORY": num(7 * per_conn, 0),
        "SFU_PER_STREAM_L1": num((val(s, 4, "downKbps") / 3 + val(s, 6, "downKbps") / 5) / 2),
        "MESH_CPU_RATIO": num(val(m, 4, "cpu") / val(m, 2, "cpu"), 1),
        "CPU_SAVING_4": num(100 * (1 - val(s, 4, "cpu") / val(m, 4, "cpu"))),
        "SFU_MAX_AJIT": num(max(r["audioJitter"]["mean"] for r in SUMMARY["labels"][s]["rows"] if r.get("audioJitter")), 2),
        "MESH_FREEZE_6": num(val(m, 6, "freezePerClient"), 0),
        "MESH_LIMITED_6": num(100 * val(m, 6, "limited"), 0),
        "SINGLE_VS_SIM_8": num(val("sfu-nosimulcast-norender", 8, "downKbps") / val(s, 8, "downKbps"), 1),
        "L0_SAVING_6": num(100 * (1 - val("sfu-layer0-norender", 6, "downKbps") / val(s, 6, "downKbps"))),
        "SFU_EGRESS_10": num(10 * val(s, 10, "downKbps") / 1000, 1),
        "TOTAL_CPU": (
            f"Mesh dùng tổng cộng khoảng {num(4 * val(m, 4, 'cpu'))}% một lõi ở N = 4, "
            f"{num(6 * val(m, 6, 'cpu'))}% ở N = 6 và {num(8 * val(m, 8, 'cpu'))}% ở N = 8, trong khi SFU dùng "
            f"{num(4 * val(s, 4, 'cpu'))}% ở N = 4 và {num(10 * val(s, 10, 'cpu'))}% ở N = 10, cộng thêm "
            f"{num(val(s, 10, 'workerCpu'), 1)}% cho worker"
        ) if LANG == "vi" else (
            f"the mesh used about {num(4 * val(m, 4, 'cpu'))}% of one core in total at N = 4, "
            f"{num(6 * val(m, 6, 'cpu'))}% at N = 6 and {num(8 * val(m, 8, 'cpu'))}% at N = 8, whereas the SFU used "
            f"{num(4 * val(s, 4, 'cpu'))}% at N = 4 and {num(10 * val(s, 10, 'cpu'))}% at N = 10, plus "
            f"{num(val(s, 10, 'workerCpu'), 1)}% for its worker"
        ),
    }
    sfu_up = [r["upKbps"]["mean"] for r in SUMMARY["labels"][s]["rows"]]
    sfu_cpu_max = max(r["cpu"]["mean"] for r in SUMMARY["labels"][s]["rows"])
    computed["O3_EVIDENCE"] = (
        f"uplink SFU {num(min(sfu_up))}–{num(max(sfu_up))} kbps từ N = 2 đến N = 10, không mất gói (20 fps từ N = 4); "
        f"Mesh sập ở N = 6"
    ) if LANG == "vi" else (
        f"SFU uplink {num(min(sfu_up))}–{num(max(sfu_up))} kbps from N = 2 to N = 10 with no packet loss (20 fps from N = 4); "
        f"the mesh baseline collapsed at N = 6"
    )
    computed["O3_SUMMARY"] = (
        f"Từ N = 2 đến N = 10, uplink của một client SFU nằm trong khoảng {num(min(sfu_up))}–{num(max(sfu_up))} kbps "
        f"và CPU dưới {num(sfu_cpu_max, 0)}% một lõi, trong khi CPU của một client Mesh tăng "
        f"{num(val(m, 4, 'cpu') / val(m, 2, 'cpu'), 1)} lần từ N = 2 lên N = 4 và phiên Mesh sập ở N = 6."
    ) if LANG == "vi" else (
        f"From N = 2 to N = 10 the uplink of an SFU client stayed between {num(min(sfu_up))} and {num(max(sfu_up))} kbps "
        f"and its CPU below {num(sfu_cpu_max, 0)}% of one core, while the CPU of a mesh client grew "
        f"{num(val(m, 4, 'cpu') / val(m, 2, 'cpu'), 1)} times from N = 2 to N = 4 and the mesh session collapsed at N = 6."
    )
    for k, v in computed.items():
        text = text.replace(f"@{k}@", v)
    tests = HERE / "build" / "data" / ("tests.md" if LANG == "en" else "tests-vi.md")
    if tests.exists():
        text = text.replace("@TESTS_TABLE@", tests.read_text(encoding="utf-8").strip())
    for leftover in re.findall(r"@[A-Z_0-9]+@|\{\{[^}]+\}\}", text):
        warnings.append(f"unfilled token {leftover}")
    return text


def media_table_md(label):
    out = ["| N | Uplink (kbps) | Downlink (kbps) | CPU (%) | Encoder | RTT (ms) | Jitter video (ms) | Mất gói (%) | FPS |"
           if LANG == "vi" else
           "| N | Uplink (kbps) | Downlink (kbps) | CPU (%) | Encoders | RTT (ms) | Video jitter (ms) | Loss (%) | FPS |",
           "| --- | --- | --- | --- | --- | --- | --- | --- | --- |"]
    for r in SUMMARY["labels"][label]["rows"]:
        n = r["clients"]
        out.append("| " + " | ".join([str(n)] + [cell(label, n, k, br=True) for k in
                   ("upKbps", "downKbps", "cpu", "encoders", "rtt", "videoJitter", "loss", "fps")]) + " |")
    return out


def server_table_md(label):
    out = ["| N | CPU worker mediasoup (% một lõi) | CPU Node.js (% một lõi) | RSS Node.js (MiB) | Tổng downlink mọi client (Mbps, tính lại) |"
           if LANG == "vi" else
           "| N | mediasoup worker CPU (% of one core) | Node.js CPU (% of one core) | Node.js RSS (MiB) | Downlink sum of all clients (Mbps, computed) |",
           "| --- | --- | --- | --- | --- |"]
    for r in SUMMARY["labels"][label]["rows"]:
        n = r["clients"]
        out.append(f"| {n} | {cell(label, n, 'workerCpu')} | {cell(label, n, 'nodeCpu')} | {cell(label, n, 'rss')} | "
                   f"{num(n * val(label, n, 'downKbps') / 1000, 1)} |")
    return out


def variant_table_md():
    out = [("| Cấu hình |" if LANG == "vi" else "| Configuration |") + " N | Uplink (kbps) | Downlink (kbps) | CPU (%) | FPS |",
           "| --- | --- | --- | --- | --- | --- |"]
    names = ({"sfu-norender": "Simulcast, chính sách mặc định", "sfu-nosimulcast-norender": "Một encoding",
              "sfu-layer0-norender": "Simulcast, ép lớp 0"} if LANG == "vi" else
             {"sfu-norender": "Simulcast, default policy", "sfu-nosimulcast-norender": "Single encoding",
              "sfu-layer0-norender": "Simulcast, layer 0 forced"})
    for label, name in names.items():
        for r in SUMMARY["labels"].get(label, {}).get("rows", []):
            n = r["clients"]
            if label == "sfu-norender" and n not in (2, 6, 8, 10):
                continue
            out.append(f"| {name} | {n} | " + " | ".join(cell(label, n, k, with_sd=False) for k in ("upKbps", "downKbps", "cpu", "fps")) + " |")
    return out


def expand_directives(text):
    lines = []
    for line in text.splitlines():
        s = line.strip()
        if s.startswith(("MEDIATABLE:", "SERVERTABLE:", "VARIANTTABLE:")):
            kind, rest = s.split(":", 1)
            if kind == "VARIANTTABLE":
                lines += [f"TABLE: {rest.strip()}"] + variant_table_md()
            else:
                label, caption = [x.strip() for x in rest.split("|", 1)]
                lines += [f"TABLE: {caption}"] + (media_table_md(label) if kind == "MEDIATABLE" else server_table_md(label))
        else:
            lines.append(line)
    return "\n".join(lines)


# ---------------------------------------------------------------- markdown
def parse_table(lines, i):
    rows = []
    while i < len(lines) and lines[i].startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def render_markdown(cur, text, new_num):
    lines = text.splitlines()
    i = 0
    list_num = None
    list_kind = None
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            list_num = None
            i += 1
            continue
        if s.startswith("# "):
            p = cur.para("Heading 1")
            add_run(p, s[2:].strip().upper())
        elif s.startswith("### "):
            p = cur.para("Heading 3")
            add_run(p, s[4:].strip())
        elif s.startswith("## "):
            p = cur.para("Heading 2")
            add_run(p, s[3:].strip())
        elif s.startswith("TABLE:"):
            rows, i = parse_table(lines, i + 1)
            table_block(cur, s[6:].strip(), rows)
            continue
        elif s.startswith("FIG:"):
            file, caption, width = [x.strip() for x in s[4:].split("|")]
            figure_block(cur, file, caption, width)
        elif s.startswith("CODE:"):
            caption = s[5:].strip()
            assert lines[i + 1].startswith("```"), f"code block expected after {caption}"
            j = i + 2
            body = []
            while not lines[j].startswith("```"):
                body.append(lines[j])
                j += 1
            code_block(cur, caption, body)
            i = j + 1
            continue
        elif re.match(r"^(- |\d+\. )", s):
            kind = "bullet" if s.startswith("- ") else "decimal"
            if list_num is None or list_kind != kind:
                list_num, list_kind = new_num(kind), kind
            p = cur.para(BODY)
            number_para(p, list_num)
            add_inline(p, re.sub(r"^(- |\d+\. )", "", s))
        else:
            list_num = None
            p = cur.para(BODY)
            add_inline(p, s)
        i += 1


# ---------------------------------------------------------------- front matter
def fill_front_matter(body_children):
    P = lambda i: docx.text.paragraph.Paragraph(body_children[i], doc._body)
    for i in (8, 37):
        set_text(P(i), f"{AUTHOR} - {STUDENT_ID}")
    for i in (12, 40):
        set_text(P(i), T["thesis"])
    for i in (16, 42):
        p = P(i)
        add_run(p, TITLE, bold=True, size=15)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.right_indent = Cm(0.5)
    # compensate the extra title lines on both covers
    for i in sorted([18, 19, 44], reverse=True):
        el = body_children[i]
        if not "".join(el.itertext()).strip():
            el.getparent().remove(el)
    if LANG == "vi":
        return ACK_VI
    ack = (
        "I would like to express my sincere gratitude to Mr. Vu Dinh Hong, my instructor and mentor, "
        "for his guidance throughout this graduation thesis. His valuable suggestions, his patience and his "
        "constructive feedback on the design, the experiments and this report greatly contributed to the "
        "completion of the project."
    )
    ack2 = (
        "I also thank the lecturers of the Faculty of Information Technology, Ton Duc Thang University, for the "
        "knowledge they shared during my studies, and my family and friends, who supported me and helped test the "
        "application. Despite my efforts, the thesis certainly still has shortcomings, and I look forward to "
        "receiving comments from the lecturers."
    )
    return ack, ack2


def main():
    global doc
    load_refs()
    doc = docx.Document(str(TEMPLATE))
    ensure_styles()
    new_num = setup_numbering()
    body = doc.element.body
    children = list(body.iterchildren())

    # --- body section: remove the placeholder content after the TOC section break (index 74)
    toc_sect_para = children[74]
    final_sect = children[-1]
    for el in children[75:-1]:
        body.remove(el)

    # --- front matter text (indices refer to the original template)
    if LANG == "vi":
        for i, text_ in TEMPLATE_VI.items():
            set_text(docx.text.paragraph.Paragraph(children[i], doc._body), text_)
        for t_el in children[73].iter(qn("w:t")):
            if t_el.text and "TABLE OF CONTENTS" in t_el.text:
                t_el.text = t_el.text.replace("TABLE OF CONTENTS", "MỤC LỤC")
        # proofing language of the whole document
        for lang_el in doc.styles.element.iter(qn("w:lang")):
            lang_el.set(qn("w:val"), "vi-VN")
    for lvl_text in doc.part.numbering_part.element.iter(qn("w:lvlText")):
        if lvl_text.get(qn("w:val")) == "CHAPTER %1. ":
            lvl_text.set(qn("w:val"), T["chapter"])
    ack, ack2 = fill_front_matter(children)
    P = lambda el: docx.text.paragraph.Paragraph(el, doc._body)
    p_ack = P(children[53])
    set_text(p_ack, ack)
    p2 = p_after(children[53], BODY)
    add_run(p2, ack2)
    for i in (55, 65):
        p = P(children[i])
        set_text(p, DATE_LINE)
        for r in p.runs:
            r.bold = False
            r.italic = True
            r.font.size = Pt(13)
    for i in (57, 67):
        p = P(children[i])
        set_text(p, T["signature"])
    for i in (58, 69):
        p = P(children[i])
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.left_indent = Cm(6.35)
        p.paragraph_format.space_before = Pt(48)
        add_run(p, T["signer"], bold=True, size=13)
    decl = DECL_VI if LANG == "vi" else [
        "I hereby declare that this graduation thesis is my own work, carried out under the guidance of "
        "Mr. Vu Dinh Hong. The research content and the results presented in it are original and have not "
        "been published in any form before. The measurements in the tables and figures used for analysis "
        "and evaluation were collected by me from the experiments described in Chapter 5, and the sources "
        "are stated clearly in the report.",
        "The thesis also uses ideas, specifications and data from other authors and organisations; all of "
        "them are cited and listed in the references.",
        "If any fraud is found, I take full responsibility for the content of my thesis. Ton Duc Thang "
        "University is not involved in any copyright violation that I may have caused in the course of "
        "this work.",
    ]
    for el, text_ in zip(children[62:65], decl):
        set_text(P(el), text_)

    # --- abstract + TOC + lists, inside the roman-numbered TOC section
    front = (CONTENT / "00-front.md").read_text(encoding="utf-8")
    abstract_md = front.split("@ABSTRACT", 1)[1].split("@ABBREVIATIONS", 1)[0].strip()
    abbr_md = front.split("@ABBREVIATIONS", 1)[1].strip()
    toc_sdt = children[73]
    cur = Cursor(children[72])
    t = cur.para("Tiêu đề các trang mở đầu")
    add_run(t, T["abstract"])
    render_markdown(cur, abstract_md, new_num)
    pb = cur.para("Normal")
    pb.add_run().add_break(WD_BREAK.PAGE)
    # after the TOC content control: abbreviations, list of figures, list of tables
    cur = Cursor(toc_sdt)
    for title, instr in (
        (T["abbr"], None),
        (T["figures"], 'TOC \\h \\z \\t "Figure Caption,9"'),
        (T["tables"], 'TOC \\h \\z \\t "Table Caption,8"'),
    ):
        t = cur.para("Tiêu đề các trang mở đầu")
        t.paragraph_format.page_break_before = True
        add_run(t, title)
        if instr:
            add_field(cur.para("Normal"), instr)
        else:
            rows, _ = parse_table(abbr_md.splitlines(), 0)
            table_block(cur, "", rows, compact=True)
    # drop the empty caption paragraph created for the abbreviation table
    for p in doc.paragraphs:
        if p.style.name == "Table Caption" and not p.text.strip():
            p._element.getparent().remove(p._element)
            break

    # --- chapters
    global SUMMARY
    import json
    SUMMARY = json.loads((HERE / "build" / "data" / "summary.json").read_text(encoding="utf-8"))
    cur = Cursor(toc_sect_para)
    for f in sorted(CONTENT.glob("0[1-6]-*.md")):
        render_markdown(cur, substitute(expand_directives(f.read_text(encoding="utf-8"))), new_num)
    # --- references
    h = cur.para("Heading 1")
    add_run(h, T["refs"])
    pPr = h._element.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    for tag, val in (("w:ilvl", "0"), ("w:numId", "0")):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        numPr.append(e)
    pPr.insert(1, numPr)
    for n, key in enumerate(ref_order, 1):
        p = cur.para("Normal")
        set_ind(p, left=567, hanging=567)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT  # long URLs would stretch justified lines
        add_run(p, f"[{n}]\t", size=12)
        add_inline(p, REFS[key], size=12)
    unused = [k for k in REFS if k not in ref_order]

    # --- figure/table reference check (faculty rule: every figure is referenced in the text)
    texts = [p.text for p in doc.paragraphs if p.style.name not in ("Figure Caption", "Table Caption", "Listing Caption")]
    texts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
    joined = "\n".join(texts)
    for p in doc.paragraphs:
        if p.style.name in ("Figure Caption", "Table Caption", "Listing Caption"):
            m = re.match(rf"((?:{T['caption_words']}) \d+\.\d+)", p.text)
            if m and not re.search(re.escape(m.group(1)) + r"(?!\d)", joined):
                warnings.append(f"not referenced in text: {m.group(1)}")

    # The TOC and the lists of figures/tables are filled by finalize.ps1 (Microsoft Word).
    doc.core_properties.author = T["signer"]
    doc.core_properties.title = TITLE.title()
    doc.save(str(OUT))
    print(f"saved {OUT}")
    print(f"references cited: {len(ref_order)}; unused: {unused}")
    for w in warnings:
        print("WARNING:", w)


if __name__ == "__main__":
    main()
